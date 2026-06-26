#!/usr/bin/env python3
"""
Importar leads de planilha Excel para o Supabase.

Lê a planilha mais recente de output/prospeccao/ ou output/campanhas/,
normaliza os telefones, e importa para o Supabase com deduplicação
por telefone_normalizado.

Uso:
    python import_leads_to_supabase.py
    python import_leads_to_supabase.py --arquivo caminho/do/arquivo.xlsx

NUNCA reseta status de leads já abordados.
NUNCA sobrescreve ultimo_contato_em ou proximo_followup_em de leads existentes.
"""

import argparse
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    from dotenv import load_dotenv
except ImportError:
    # Fallback: carregar .env manualmente
    def load_dotenv(path=None):
        env_path = Path(path) if path else Path(__file__).parent / ".env"
        if env_path.exists():
            with open(env_path, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith("#") and "=" in line:
                        key, _, value = line.partition("=")
                        os.environ.setdefault(key.strip(), value.strip())

try:
    from supabase import create_client
    HAS_SUPABASE = True
except ImportError:
    HAS_SUPABASE = False

try:
    from openpyxl import load_workbook
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    from openpyxl import load_workbook

# Adicionar raiz do projeto ao path
sys.path.insert(0, str(Path(__file__).parent))
from utils.phone_utils import normalizar_telefone_br, gerar_link_whatsapp
from config.regioes import resolve_regiao, get_output_dir

BASE_DIR = Path(__file__).parent
OUTPUT_DIR = BASE_DIR / "output"
PROSPECCAO_DIR = OUTPUT_DIR / "prospeccao"
CAMPANHAS_DIR = OUTPUT_DIR / "campanhas"

# Status que NUNCA devem ser resetados
STATUS_PROTEGIDOS = {"abordado", "respondeu", "follow_up", "interessado", "convertido", "perdido"}

# Mapeamento de colunas do Excel para campos do Supabase
# Aceita variações de nome encontradas nos diferentes scripts
MAPEAMENTO_COLUNAS = {
    "Nome": "nome",
    "nome": "nome",
    "name": "nome",
    "Nicho": "nicho",
    "nicho": "nicho",
    "Categoria": "categoria",
    "categoria": "categoria",
    "Cidade": "cidade",
    "cidade": "cidade",
    "Bairro": "bairro",
    "bairro": "bairro",
    "Telefone": "telefone",
    "telefone": "telefone",
    "phone": "telefone",
    "WhatsApp": "whatsapp",
    "whatsapp": "whatsapp",
    "Instagram": "instagram",
    "instagram": "instagram",
    "Email": "email",
    "email": "email",
    "Site": "tem_site",
    "tem_site": "tem_site",
    "Tem Site?": "tem_site",
    "Tem Site": "tem_site",
    "URL Site": "url_site",
    "URL do Site": "url_site",
    "url_site": "url_site",
    "Nota": "avaliacao",
    "Avaliação": "avaliacao",
    "Nota Google": "avaliacao",
    "nota_google": "avaliacao",
    "avaliacao": "avaliacao",
    "Avaliações": "num_avaliacoes",
    "Nº Avaliações": "num_avaliacoes",
    "N. Avaliações": "num_avaliacoes",
    "num_avaliacoes": "num_avaliacoes",
    "qtd_avaliacoes": "num_avaliacoes",
    "Score": "score",
    "score": "score",
    "Prioridade": "prioridade",
    "prioridade": "prioridade",
    "Oferta Sugerida": "oferta_sugerida",
    "oferta_sugerida": "oferta_sugerida",
    "Motivo Prioridade": "observacoes",
    "Motivo": "observacoes",
    "Mensagem WhatsApp": "mensagem_whatsapp",
    "mensagem_whatsapp": "mensagem_whatsapp",
    "Link WhatsApp": "link_whatsapp",
    "link_whatsapp": "link_whatsapp",
    "Endereço": "endereco",
    "Endereco": "endereco",
    "endereco": "endereco",
    "Link Maps": "link_maps",
    "Ação Recomendada": "acao_recomendada",
    "Tipo de Material": "tipo_de_material",
    "Observações": "observacoes_extra",
    "Observacoes": "observacoes_extra",
    # Campos do modo AVGESTAO (presentes somente em planilhas avgestao)
    "Score AVGESTÃO": "score_avgestao",
    "Score AVGESTAO": "score_avgestao",
    "Faz assistência": "faz_assistencia",
    "Faz assistencia": "faz_assistencia",
    "Subnicho": "subnicho",
    "Grupo": "grupo",
    "Motivos do score": "motivos_score",
    "Nome curto": "nome_curto",
    "Mensagem inicial": "mensagem_whatsapp",
}


def encontrar_planilha(caminho_arquivo=None):
    """Encontra a planilha mais recente."""
    if caminho_arquivo:
        p = Path(caminho_arquivo)
        if p.exists():
            return p, "custom"
        print(f"❌ Arquivo não encontrado: {caminho_arquivo}")
        sys.exit(1)

    # Prioridade: campanha_diaria > leads_prospeccao > qualquer xlsx
    campanha = CAMPANHAS_DIR / "campanha_diaria.xlsx"
    if campanha.exists():
        return campanha, "campanha"

    prospeccao = PROSPECCAO_DIR / "leads_prospeccao.xlsx"
    if prospeccao.exists():
        return prospeccao, "prospeccao"

    # Fallback: procurar qualquer xlsx
    for d in [CAMPANHAS_DIR, PROSPECCAO_DIR]:
        xlsx_files = sorted(d.glob("*.xlsx"), key=lambda f: f.stat().st_mtime, reverse=True)
        if xlsx_files:
            return xlsx_files[0], "fallback"

    print("❌ Nenhuma planilha encontrada em output/prospeccao/ ou output/campanhas/")
    sys.exit(1)


def ler_planilha(caminho, tipo="campanha"):
    """Lê planilha Excel e retorna lista de dicts."""
    wb = load_workbook(str(caminho), read_only=True, data_only=True)

    # Escolher qual aba ler
    if tipo == "campanha" and "Top 50" in wb.sheetnames:
        sheets = ["Top 50"]
    elif tipo == "prospeccao" and "Leads" in wb.sheetnames:
        sheets = ["Leads"]
    else:
        # Ler todas as abas exceto "Resumo"
        sheets = [s for s in wb.sheetnames if s.lower() != "resumo"]

    todos_leads = []
    vistos = set()

    for nome_aba in sheets:
        ws = wb[nome_aba]

        # Mapear cabeçalhos
        cabecalhos = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
        col_map = {}
        for idx, cab in enumerate(cabecalhos):
            campo = MAPEAMENTO_COLUNAS.get(cab, MAPEAMENTO_COLUNAS.get(cab.lower(), ""))
            if campo:
                col_map[campo] = idx

        if not col_map:
            continue

        for row in ws.iter_rows(min_row=2, values_only=True):
            lead = {}
            for campo, idx in col_map.items():
                valor = row[idx] if idx < len(row) else None
                lead[campo] = str(valor).strip() if valor is not None else ""

            # Deduplicar por nome+cidade
            chave = (lead.get("nome", "").lower().strip(), lead.get("cidade", "").lower().strip())
            if lead.get("nome") and chave not in vistos:
                vistos.add(chave)
                todos_leads.append(lead)

    wb.close()
    return todos_leads


def normalizar_lead(lead_bruto):
    """Normaliza campos do lead para o formato do Supabase."""
    lead = {}

    # Campos diretos
    lead["nome"] = lead_bruto.get("nome", "").strip()
    lead["telefone"] = lead_bruto.get("telefone", "").strip()
    lead["whatsapp"] = lead_bruto.get("whatsapp", "").strip()
    lead["instagram"] = lead_bruto.get("instagram", "").strip()
    lead["email"] = lead_bruto.get("email", "").strip()
    lead["categoria"] = lead_bruto.get("categoria", "").strip() or lead_bruto.get("nicho", "").strip()
    lead["nicho"] = lead_bruto.get("nicho", "").strip() or lead_bruto.get("categoria", "").strip()
    lead["cidade"] = lead_bruto.get("cidade", "").strip()
    lead["bairro"] = lead_bruto.get("bairro", "").strip()
    lead["endereco"] = lead_bruto.get("endereco", "").strip()
    lead["url_site"] = lead_bruto.get("url_site", "").strip()
    lead["prioridade"] = lead_bruto.get("prioridade", "").strip()
    lead["oferta_sugerida"] = lead_bruto.get("oferta_sugerida", "").strip()
    lead["mensagem_whatsapp"] = lead_bruto.get("mensagem_whatsapp", "").strip()
    lead["link_whatsapp"] = lead_bruto.get("link_whatsapp", "").strip()

    # Telefone normalizado (chave de dedup)
    tel_raw = lead["whatsapp"] or lead["telefone"]
    lead["telefone_normalizado"] = normalizar_telefone_br(tel_raw)

    # Se link_whatsapp vazio mas temos telefone normalizado e mensagem, gerar link
    if not lead["link_whatsapp"] and lead["telefone_normalizado"] and lead["mensagem_whatsapp"]:
        lead["link_whatsapp"] = gerar_link_whatsapp(lead["telefone_normalizado"], lead["mensagem_whatsapp"])
    elif not lead["link_whatsapp"] and lead["telefone_normalizado"]:
        lead["link_whatsapp"] = gerar_link_whatsapp(lead["telefone_normalizado"])

    # tem_site: converter string para boolean
    site_val = lead_bruto.get("tem_site", "").strip().lower()
    lead["tem_site"] = site_val in ("true", "sim", "1", "yes", "verdadeiro", "s")

    # avaliacao: converter para número
    nota = lead_bruto.get("avaliacao", "0").strip()
    try:
        lead["avaliacao"] = float(nota.replace(",", ".")) if nota else 0.0
    except (ValueError, TypeError):
        lead["avaliacao"] = 0.0

    # num_avaliacoes: converter para inteiro
    qtd = lead_bruto.get("num_avaliacoes", "0").strip()
    try:
        lead["num_avaliacoes"] = int(re.sub(r"\D", "", qtd)) if qtd else 0
    except (ValueError, TypeError):
        lead["num_avaliacoes"] = 0

    # score: converter para inteiro
    score = lead_bruto.get("score", "0").strip()
    try:
        lead["score"] = int(float(score)) if score else 0
    except (ValueError, TypeError):
        lead["score"] = 0

    # Observações extras (motivo, observações)
    obs_parts = []
    if lead_bruto.get("observacoes"):
        obs_parts.append(lead_bruto["observacoes"])
    if lead_bruto.get("observacoes_extra"):
        obs_parts.append(lead_bruto["observacoes_extra"])
    lead["observacoes"] = " | ".join(filter(None, obs_parts))

    # Campos do modo AVGESTAO (opcionais - presentes somente em planilhas avgestao)
    score_avg = str(lead_bruto.get("score_avgestao", "")).strip()
    try:
        lead["score_avgestao"] = int(float(score_avg)) if score_avg else 0
    except (ValueError, TypeError):
        lead["score_avgestao"] = 0
    lead["faz_assistencia"] = lead_bruto.get("faz_assistencia", "").strip()
    lead["subnicho"] = lead_bruto.get("subnicho", "").strip()
    lead["grupo"] = lead_bruto.get("grupo", "").strip()
    lead["motivos_score"] = lead_bruto.get("motivos_score", "").strip()
    lead["nome_curto"] = lead_bruto.get("nome_curto", "").strip()

    # Detecta produto: avgestao se houver grupo ou score_avgestao informado
    lead["produto"] = "avgestao" if (lead["grupo"] or score_avg) else "landing"

    return lead


def main():
    parser = argparse.ArgumentParser(description="Importar leads de planilha para o Supabase")
    parser.add_argument("--arquivo", help="Caminho para arquivo .xlsx específico", default=None)
    parser.add_argument("--regiao", choices=["baixada", "rio_premium", "todas"], default=None,
                        help="Regiao de prospeccao (padrao: baixada)")
    args = parser.parse_args()

    print("=" * 60)
    print("📦 Importando leads para o Supabase")
    print("=" * 60)

    # Verificar dependências
    if not HAS_SUPABASE:
        print("❌ Pacote 'supabase' não instalado. Instale com: pip install supabase")
        sys.exit(1)

    # Carregar .env
    load_dotenv()
    supabase_url = os.environ.get("SUPABASE_URL", "")
    supabase_key = os.environ.get("SUPABASE_ANON_KEY") or os.environ.get("SUPABASE_KEY", "")

    if not supabase_url or not supabase_key:
        print("❌ SUPABASE_URL e SUPABASE_ANON_KEY não encontrados no .env")
        sys.exit(1)

    # Conectar ao Supabase
    try:
        sb = create_client(supabase_url, supabase_key)
        print(f"✅ Conectado ao Supabase: {supabase_url}")
    except Exception as e:
        print(f"❌ Erro ao conectar ao Supabase: {e}")
        sys.exit(1)

    # Resolver regiao
    regioes = resolve_regiao(args.regiao)
    regiao = regioes[0]  # Primeira regiao (unica, exceto --regiao todas)
    print(f"📍 Regiao: {regiao.label}")

    # Encontrar planilha (priorizar diretorio da regiao se --regiao informado)
    if not args.arquivo and args.regiao is not None:
        # Buscar planilha no diretorio especifico da regiao
        regiao_prospeccao = get_output_dir(regiao.key, "prospeccao")
        regiao_campanhas = get_output_dir(regiao.key, "campanhas")
        campanha = regiao_campanhas / "campanha_diaria.xlsx"
        if campanha.exists():
            planilha, tipo = campanha, "campanha"
        else:
            prospeccao = regiao_prospeccao / "leads_prospeccao.xlsx"
            if prospeccao.exists():
                planilha, tipo = prospeccao, "prospeccao"
            else:
                # Fallback: procurar qualquer xlsx nos dirs da regiao
                for d in [regiao_campanhas, regiao_prospeccao]:
                    xlsx_files = sorted(d.glob("*.xlsx"), key=lambda f: f.stat().st_mtime, reverse=True)
                    if xlsx_files:
                        planilha, tipo = xlsx_files[0], "fallback"
                        break
                else:
                    print(f"❌ Nenhuma planilha encontrada em {regiao_prospeccao}/ ou {regiao_campanhas}/")
                    sys.exit(1)
    else:
        planilha, tipo = encontrar_planilha(args.arquivo)
    print(f"📊 Planilha: {planilha.name} (tipo: {tipo})")

    # Ler leads
    leads_brutos = ler_planilha(planilha, tipo)
    print(f"📈 Total de linhas lidas: {len(leads_brutos)}")

    # Normalizar leads
    leads = []
    ignorados_sem_tel = 0
    for lb in leads_brutos:
        lead = normalizar_lead(lb)
        if not lead["nome"]:
            continue
        if not lead["telefone_normalizado"]:
            ignorados_sem_tel += 1
            continue
        leads.append(lead)

    print(f"📋 Leads válidos com telefone: {len(leads)}")
    if ignorados_sem_tel:
        print(f"⚠️  Leads ignorados (sem telefone válido): {ignorados_sem_tel}")

    if not leads:
        print("❌ Nenhum lead válido para importar!")
        sys.exit(1)

    # Buscar leads existentes no Supabase para deduplicação
    telefones = [l["telefone_normalizado"] for l in leads]
    existentes = {}
    batch_size = 100

    for i in range(0, len(telefones), batch_size):
        batch = telefones[i:i + batch_size]
        try:
            result = sb.table("leads").select("id,telefone_normalizado,status").in_("telefone_normalizado", batch).execute()
            for row in result.data:
                existentes[row["telefone_normalizado"]] = row
        except Exception as e:
            print(f"⚠️  Erro ao buscar leads existentes (batch {i}): {e}")

    print(f"🔍 Leads já existentes no Supabase: {len(existentes)}")

    # Importar leads
    novos = 0
    atualizados = 0
    preservados = 0
    erros = 0

    for lead in leads:
        tel = lead["telefone_normalizado"]
        existente = existentes.get(tel)

        if not existente:
            # Inserir novo lead
            novo_lead = {
                "nome": lead["nome"],
                "telefone": lead["telefone"],
                "whatsapp": lead["whatsapp"],
                "telefone_normalizado": tel,
                "instagram": lead["instagram"],
                "email": lead["email"],
                "categoria": lead["categoria"],
                "nicho": lead["nicho"],
                "cidade": lead["cidade"],
                "bairro": lead["bairro"],
                "endereco": lead["endereco"],
                "tem_site": lead["tem_site"],
                "url_site": lead["url_site"],
                "avaliacao": lead["avaliacao"],
                "num_avaliacoes": lead["num_avaliacoes"],
                "score": lead["score"],
                "prioridade": lead["prioridade"],
                "oferta_sugerida": lead["oferta_sugerida"],
                "mensagem_whatsapp": lead["mensagem_whatsapp"],
                "link_whatsapp": lead["link_whatsapp"],
                "status": "novo",
                "origem": regiao.key,
                "observacoes": lead["observacoes"],
                "produto": lead.get("produto", "landing"),
                "grupo": lead.get("grupo", ""),
                "subnicho": lead.get("subnicho", ""),
                "faz_assistencia": lead.get("faz_assistencia", ""),
                "score_avgestao": lead.get("score_avgestao", 0),
                "motivos_score": lead.get("motivos_score", ""),
                "nome_curto": lead.get("nome_curto", ""),
            }

            try:
                result = sb.table("leads").insert(novo_lead).execute()
                novo_id = result.data[0]["id"]

                # Criar interação de importação
                sb.table("lead_interactions").insert({
                    "lead_id": novo_id,
                    "tipo": "importacao",
                    "canal": "planilha",
                    "observacao": f"Importado de {planilha.name}",
                }).execute()

                novos += 1
            except Exception as e:
                print(f"  ❌ Erro ao inserir {lead['nome']}: {e}")
                erros += 1
        else:
            # Atualizar lead existente (apenas campos de dados, NUNCA status protegido)
            status_atual = existente.get("status", "novo")

            if status_atual in STATUS_PROTEGIDOS:
                preservados += 1
                # Atualizar apenas campos informativos, preservando status e datas
                update_data = {
                    "score": lead["score"],
                    "prioridade": lead["prioridade"],
                    "oferta_sugerida": lead["oferta_sugerida"],
                    "mensagem_whatsapp": lead["mensagem_whatsapp"],
                    "link_whatsapp": lead["link_whatsapp"],
                    "categoria": lead["categoria"],
                    "nicho": lead["nicho"],
                    "avaliacao": lead["avaliacao"],
                    "num_avaliacoes": lead["num_avaliacoes"],
                    "produto": lead.get("produto", "landing"),
                    "grupo": lead.get("grupo", ""),
                    "subnicho": lead.get("subnicho", ""),
                    "faz_assistencia": lead.get("faz_assistencia", ""),
                    "score_avgestao": lead.get("score_avgestao", 0),
                    "motivos_score": lead.get("motivos_score", ""),
                    "nome_curto": lead.get("nome_curto", ""),
                }
            else:
                atualizados += 1
                # Atualizar dados e permitir mudança de status se for 'novo' -> 'pronto_para_enviar'
                novo_status = status_atual
                if status_atual == "novo" and lead["score"] >= 65:
                    novo_status = "pronto_para_enviar"

                update_data = {
                    "nome": lead["nome"],
                    "score": lead["score"],
                    "prioridade": lead["prioridade"],
                    "oferta_sugerida": lead["oferta_sugerida"],
                    "mensagem_whatsapp": lead["mensagem_whatsapp"],
                    "link_whatsapp": lead["link_whatsapp"],
                    "categoria": lead["categoria"],
                    "nicho": lead["nicho"],
                    "telefone": lead["telefone"],
                    "whatsapp": lead["whatsapp"],
                    "tem_site": lead["tem_site"],
                    "url_site": lead["url_site"],
                    "avaliacao": lead["avaliacao"],
                    "num_avaliacoes": lead["num_avaliacoes"],
                    "status": novo_status,
                    "bairro": lead["bairro"],
                    "endereco": lead["endereco"],
                    "produto": lead.get("produto", "landing"),
                    "grupo": lead.get("grupo", ""),
                    "subnicho": lead.get("subnicho", ""),
                    "faz_assistencia": lead.get("faz_assistencia", ""),
                    "score_avgestao": lead.get("score_avgestao", 0),
                    "motivos_score": lead.get("motivos_score", ""),
                    "nome_curto": lead.get("nome_curto", ""),
                }

            try:
                sb.table("leads").update(update_data).eq("id", existente["id"]).execute()
            except Exception as e:
                print(f"  ❌ Erro ao atualizar {lead['nome']}: {e}")
                erros += 1

    # Criar registro de campanha
    try:
        sb.table("campaigns").insert({
            "nome": f"Importação {planilha.name}",
            "data": date.today().isoformat(),
            "quantidade_leads": len(leads),
            "observacao": f"Tipo: {tipo} | Novos: {novos} | Atualizados: {atualizados} | Preservados: {preservados}",
        }).execute()
    except Exception:
        pass  # Campanha é informativa, não crítica

    # Resumo
    print()
    print("=" * 60)
    print("📊 RESUMO DA IMPORTAÇÃO")
    print("=" * 60)
    print(f"  📄 Leads lidos da planilha:  {len(leads_brutos)}")
    print(f"  ✅ Leads válidos com tel:     {len(leads)}")
    print(f"  🆕 Novos inseridos:          {novos}")
    print(f"  🔄 Atualizados:              {atualizados}")
    print(f"  🔒 Preservados (status ok):  {preservados}")
    print(f"  ⚠️  Ignorados (sem tel):     {ignorados_sem_tel}")
    print(f"  ❌ Erros:                    {erros}")
    print("=" * 60)

    return novos + atualizados


if __name__ == "__main__":
    main()