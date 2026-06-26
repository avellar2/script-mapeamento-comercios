"""
Escrita do XLSX do modo AVGESTAO (compartilhado).

Usa o spec de colunas de config.avgestao.COLUNAS_XLSX_AVGESTAO e aplica
formatacao basica (cabecalho, larguras, filtros, freeze, cores por score).
Nao envia nada automaticamente - apenas gera o arquivo.
"""

from __future__ import annotations

from pathlib import Path

from config.avgestao import COLUNAS_XLSX_AVGESTAO, linha_xlsx_avgestao, linha_xlsx_avgestao_pronto, prioridade_avgestao


def exportar_xlsx_avgestao(leads: list[dict], caminho: str | Path, titulo_aba: str = "Leads", enriquecer: bool = True) -> Path:
    """
    Exporta leads para XLSX no formato AVGESTAO.

    enriquecer=True  -> recalcula score/mensagem (uso: prospectar)
    enriquecer=False -> preserva score/mensagem ja existentes (uso: campanha)

    NUNCA envia mensagens - apenas gera o arquivo com link wa.me preenchido.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = titulo_aba[:31]

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    for col_idx, (nome, _campo, largura) in enumerate(COLUNAS_XLSX_AVGESTAO, 1):
        cell = ws.cell(row=1, column=col_idx, value=nome)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin
        ws.column_dimensions[get_column_letter(col_idx)].width = largura

    cor_alta = PatternFill("solid", fgColor="ECFDF5")
    cor_media = PatternFill("solid", fgColor="FFFBEB")
    cor_baixa = PatternFill("solid", fgColor="FEF2F2")
    font_link = Font(color="2563EB", underline="single")

    builder = linha_xlsx_avgestao if enriquecer else linha_xlsx_avgestao_pronto
    linhas = [builder(l) for l in leads]

    for row_idx, linha in enumerate(linhas, 2):
        try:
            score = int(float(str(linha.get("Score AVGESTÃO", 0))))
        except (ValueError, TypeError):
            score = 0
        prioridade = prioridade_avgestao(score)
        row_fill = cor_alta if prioridade == "Alta" else cor_media if prioridade == "Média" else cor_baixa

        for col_idx, (header, _campo, _largura) in enumerate(COLUNAS_XLSX_AVGESTAO, 1):
            valor = linha.get(header, "")
            if header == "Score AVGESTÃO":
                valor = score
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=(col_idx >= 13))
            cell.fill = row_fill

            if header == "Link WhatsApp" and str(valor).startswith("https://wa.me"):
                cell.font = font_link
                cell.hyperlink = valor

            if header == "Score AVGESTÃO":
                if prioridade == "Alta":
                    cell.font = Font(bold=True, color="059669")
                elif prioridade == "Média":
                    cell.font = Font(bold=True, color="D97706")
                else:
                    cell.font = Font(bold=True, color="DC2626")

    if linhas:
        last_col = get_column_letter(len(COLUNAS_XLSX_AVGESTAO))
        ws.auto_filter.ref = f"A1:{last_col}{len(linhas) + 1}"
    ws.freeze_panes = "A2"

    wb.save(caminho)
    return caminho
