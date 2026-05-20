"""
Painel de Prospecção - Gerador de Dashboard HTML
=================================================
Lê a planilha mais atual e gera um painel visual em HTML
para controlar a prospecção diária de vendas de landing pages.

Uso: python gerar_painel_prospeccao.py [--limite N] [--novos]
"""

import json
import re
import sqlite3
import sys
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import quote

try:
    from openpyxl import load_workbook
except ImportError:
    print("[!] openpyxl não instalado. Instalando...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    from openpyxl import load_workbook


# ═══════════════════════════════════════════════════════════════════
# CONFIGURAÇÃO
# ═══════════════════════════════════════════════════════════════════

DAILY_LIMIT = 20

BASE_DIR = Path(__file__).parent
OUTPUT_DIR = BASE_DIR / "output"
PAINEL_DIR = OUTPUT_DIR / "painel"
PROSPECCAO_DIR = OUTPUT_DIR / "prospeccao"
CAMPANHAS_DIR = OUTPUT_DIR / "campanhas"

# ═══════════════════════════════════════════════════════════════════
# CLASSIFICAÇÃO POR NICHO
# ═══════════════════════════════════════════════════════════════════

NICHOS_ESTETICA = [
    "estética", "estetica", "salão de beleza", "salao de beleza",
    "manicure", "sobrancelha", "cílios", "cilios", "depilação",
    "depilacao", "spa", "clínica estética", "clinica estetica",
    "beauty", "nail", "cabeleireiro", "cabeleireira", "hair",
    "maquiagem", "maquiador", "maquiadora",
]

NICHOS_BARBEARIA = [
    "barbearia", "barbeiro", "salão masculino", "salao masculino",
    "barber", "barbershop",
]

NICHOS_COMIDA = [
    "restaurante", "marmitaria", "pizzaria", "lanchonete", "açaí",
    "açai", "confeitaria", "padaria", "churrascaria", "bar",
    "comida", "delivery", "hambúrguer", "hamburguer", "hot dog",
    "sushi", "japonesa", "acarajé", "acaraje", "petisco",
    "cardápio", "cardapio", "doces", "bolo", "confeitaria",
]

NICHOS_DENTISTA = [
    "dentista", "clínica odontológica", "clinica odontologica",
    "odontologia", "ortodontia", "implante", "dentística",
    "clínica médica", "clinica medica", "médico", "medico",
    "consultório", "consultorio", "saúde", "saude",
]

NICHOS_ADVOCACIA = [
    "advogado", "advogada", "advocacia", "escritório de advocacia",
    "escritorio de advocacia", "jurídico", "juridico", "direito",
    "advocacia", "defensoria",
]

NICHOS_SERVICOS = [
    "assistência técnica", "assistencia tecnica", "refrigeração",
    "refrigeracao", "eletricista", "encanador", "manutenção",
    "manutencao", "conserto", "serralheria", "marcenaria", "pintor",
    "vidraçaria", "vidracaria", "informática", "informatica",
    "automação", "automacao", "reparo", "técnico", "tecnico",
]

NICHOS_AUTOMOTIVA = [
    "oficina mecânica", "oficina mecanica", "automotiva", "automotivo",
    "auto", "carro", "mecânica", "mecanica", "motor", "pneu",
    "borracharia", "lava", "lavagem", "estética automotiva",
    "funilaria", "pintura automotiva",
]

NICHOS_IGREJA = [
    "igreja", "evento", "ministério", "ministerio", "culto",
    "comunidade", "paróquia", "paroquia", "capela",
]

NICHOS_CURSO = [
    "curso", "reforço", "reforco escolar", "escola", "aula",
    "professor", "professora", "ensino", "idioma", "inglês",
    "ingles", "espanhol", "pré-vestibular", "pre-vestibular",
    "concursos", "concurso", "treinamento", "capacitação",
    "capacitacao", "auto escola", "autoescola", "cfc",
]

NICHOS_PERSONALIZADOS = [
    "personalizados", "personalizado", "doces", "festas", "festa",
    "bolo decorado", "salgadinho", "convite", "lembrança",
    "lembranca", "decoração", "decoracao", "artesanal", "artesanato",
    "gift", "kit", "cesta",
]

NICHOS_BONS = (
    NICHOS_ESTETICA + NICHOS_BARBEARIA + NICHOS_COMIDA +
    NICHOS_DENTISTA + NICHOS_ADVOCACIA + NICHOS_SERVICOS +
    NICHOS_AUTOMOTIVA + NICHOS_IGREJA + NICHOS_CURSO +
    NICHOS_PERSONALIZADOS
)

# ═══════════════════════════════════════════════════════════════════
# MAPEAMENTO DE COLUNAS
# ═══════════════════════════════════════════════════════════════════

MAPEAMENTO_EXCEL = {
    "Nome": "nome",
    "Nicho": "nicho",
    "Categoria": "nicho",
    "Cidade": "cidade",
    "Bairro": "bairro",
    "Telefone": "telefone",
    "WhatsApp": "whatsapp",
    "Instagram": "instagram",
    "Site": "site",
    "Tem Site?": "site",
    "URL Site": "url_site",
    "URL do Site": "url_site",
    "Nota": "nota_google",
    "Avaliação": "nota_google",
    "Nota Google": "nota_google",
    "Avaliações": "qtd_avaliacoes",
    "Nº Avaliações": "qtd_avaliacoes",
    "Score": "score",
    "Prioridade": "prioridade",
    "Oferta Sugerida": "oferta_sugerida",
    "Motivo Prioridade": "motivo_prioridade",
    "Mensagem WhatsApp": "mensagem_whatsapp",
    "Link WhatsApp": "link_whatsapp",
    "Endereço": "endereco",
    "Endereco": "endereco",
    "Email": "email",
    "Link Maps": "link_maps",
    "Ação Recomendada": "acao_recomendada",
    "Tipo de Material": "tipo_de_material",
    "Status": "status",
    "Data Abordagem": "data_abordagem",
    "Data Follow-up": "data_followup",
    "Observações": "observacoes",
}


# ═══════════════════════════════════════════════════════════════════
# LEITURA DO EXCEL
# ═══════════════════════════════════════════════════════════════════

def ler_excel(caminho, aba="Leads"):
    """Lê planilha Excel e retorna lista de dicts."""
    wb = load_workbook(caminho, read_only=True, data_only=True)

    # Tenta a aba especificada, senão usa a primeira
    if aba in wb.sheetnames:
        ws = wb[aba]
    else:
        ws = wb[wb.sheetnames[0]]

    cabecalhos = []
    for cell in ws[1]:
        cabecalhos.append(str(cell.value).strip() if cell.value else "")

    col_map = {}
    for idx, cab in enumerate(cabecalhos):
        nome_interno = MAPEAMENTO_EXCEL.get(cab, cab.lower().strip())
        col_map[nome_interno] = idx

    leads = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        lead = {}
        for nome_interno, idx in col_map.items():
            if idx < len(row):
                lead[nome_interno] = str(row[idx]) if row[idx] is not None else ""
            else:
                lead[nome_interno] = ""
        if lead.get("nome"):
            leads.append(lead)

    wb.close()
    return leads


def ler_todas_abas(caminho):
    """Lê todas as abas de dados (não Resumo) e retorna leads únicos."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    todos = []
    vistos = set()

    for nome_aba in wb.sheetnames:
        if nome_aba.lower() == "resumo":
            continue
        ws = wb[nome_aba]

        cabecalhos = []
        for cell in ws[1]:
            cabecalhos.append(str(cell.value).strip() if cell.value else "")

        col_map = {}
        for idx, cab in enumerate(cabecalhos):
            nome_interno = MAPEAMENTO_EXCEL.get(cab, cab.lower().strip())
            col_map[nome_interno] = idx

        for row in ws.iter_rows(min_row=2, values_only=True):
            lead = {}
            for nome_interno, idx in col_map.items():
                if idx < len(row):
                    lead[nome_interno] = str(row[idx]) if row[idx] is not None else ""
                else:
                    lead[nome_interno] = ""
            if lead.get("nome"):
                chave = (lead.get("nome", "").strip().lower(), lead.get("cidade", "").strip().lower())
                if chave not in vistos:
                    vistos.add(chave)
                    todos.append(lead)

    wb.close()
    return todos


def encontrar_planilha():
    """Encontra a planilha mais atual na ordem: campanha_diaria > leads_prospeccao."""
    campanha = CAMPANHAS_DIR / "campanha_diaria.xlsx"
    if campanha.exists():
        return campanha, "campanha"

    prospeccao = PROSPECCAO_DIR / "leads_prospeccao.xlsx"
    if prospeccao.exists():
        return prospeccao, "prospeccao"

    # Fallback: busca qualquer xlsx
    for d in [CAMPANHAS_DIR, PROSPECCAO_DIR]:
        xlsx_files = list(d.glob("*.xlsx"))
        if xlsx_files:
            return xlsx_files[0], "fallback"

    return None, None


# ═══════════════════════════════════════════════════════════════════
# NORMALIZAÇÃO
# ═══════════════════════════════════════════════════════════════════

def normalizar_lead(lead):
    """Normaliza campos do lead."""
    n = {}
    n["nome"] = str(lead.get("nome", "")).strip()
    n["nicho"] = str(lead.get("nicho", lead.get("categoria", ""))).strip()
    n["cidade"] = str(lead.get("cidade", "")).strip()
    n["bairro"] = str(lead.get("bairro", "")).strip()
    n["telefone"] = str(lead.get("telefone", "")).strip()
    n["whatsapp"] = str(lead.get("whatsapp", "")).strip()

    # Site
    site_val = str(lead.get("site", "")).strip().lower()
    if site_val in ("true", "sim", "1", "yes"):
        n["site"] = "Sim"
    else:
        n["site"] = "Não"

    n["url_site"] = str(lead.get("url_site", "")).strip()

    # Nota Google
    nota = str(lead.get("nota_google", lead.get("avaliacao", "0"))).strip()
    try:
        n["nota_google"] = float(nota.replace(",", "."))
    except (ValueError, TypeError):
        n["nota_google"] = 0

    # Qtd avaliações
    qtd = str(lead.get("qtd_avaliacoes", lead.get("num_avaliacoes", "0"))).strip()
    try:
        n["qtd_avaliacoes"] = int(re.sub(r"\D", "", qtd))
    except (ValueError, TypeError):
        n["qtd_avaliacoes"] = 0

    # Score
    try:
        n["score"] = int(lead.get("score", 0))
    except (ValueError, TypeError):
        n["score"] = 0

    n["prioridade"] = str(lead.get("prioridade", "")).strip()
    n["oferta_sugerida"] = str(lead.get("oferta_sugerida", "")).strip()
    n["motivo_prioridade"] = str(lead.get("motivo_prioridade", "")).strip()
    n["acao_recomendada"] = str(lead.get("acao_recomendada", "")).strip()
    n["tipo_de_material"] = str(lead.get("tipo_de_material", "")).strip()
    n["mensagem_whatsapp"] = str(lead.get("mensagem_whatsapp", "")).strip()
    n["link_whatsapp"] = str(lead.get("link_whatsapp", "")).strip()
    n["instagram"] = str(lead.get("instagram", "")).strip()
    n["email"] = str(lead.get("email", "")).strip()
    n["endereco"] = str(lead.get("endereco", "")).strip()
    n["link_maps"] = str(lead.get("link_maps", "")).strip()

    # Status
    status_val = str(lead.get("status", "")).strip().lower()
    status_validos = [
        "novo", "abordar hoje", "mensagem enviada", "respondeu",
        "video enviado", "interessado", "proposta enviada",
        "fechado", "perdido", "follow-up", "não abordar",
    ]
    n["status"] = status_val if status_val in status_validos else "novo"
    n["data_abordagem"] = str(lead.get("data_abordagem", "")).strip()
    n["data_followup"] = str(lead.get("data_followup", lead.get("data_follow-up", ""))).strip()
    n["observacoes"] = str(lead.get("observacoes", "")).strip()

    # ID único para rastreamento
    n["lead_id"] = gerar_id(n["nome"], n["cidade"])

    # Gerar mensagem se não existir
    if not n["mensagem_whatsapp"]:
        n["mensagem_whatsapp"] = gerar_mensagem(n)

    # Gerar link WhatsApp se não existir
    if not n["link_whatsapp"] and (n["whatsapp"] or n["telefone"]):
        n["link_whatsapp"] = gerar_link_whatsapp(n)

    return n


def gerar_id(nome, cidade):
    """Gera ID único para o lead."""
    base = f"{nome}|{cidade}".lower().strip()
    return re.sub(r"[^a-z0-9]+", "-", base)[:60]


# ═══════════════════════════════════════════════════════════════════
# MENSAGENS WHATSAPP POR NICHO
# ═══════════════════════════════════════════════════════════════════

def gerar_mensagem(lead):
    """Gera mensagem WhatsApp personalizada por nicho."""
    nome = lead.get("nome", "").strip()
    nicho = lead.get("nicho", "").strip().lower()

    # Estética / Beleza
    if any(n in nicho for n in NICHOS_ESTETICA):
        return (
            f"Oi, tudo bem? Vi a {nome} no Google e percebi que vocês têm boas "
            f"avaliações. Notei que uma página simples com serviços, fotos, localização "
            f"e botão direto para agendamento no WhatsApp poderia deixar tudo mais "
            f"organizado para quem chega pelo Instagram ou Google. Quando o cliente "
            f"precisa procurar muito, ele pode acabar desistindo ou chamando outro lugar. "
            f"Eu montei uma prévia visual de como uma página assim poderia ficar para "
            f"um negócio de estética. Posso te mandar?"
        )

    # Barbearia
    if any(n in nicho for n in NICHOS_BARBEARIA):
        return (
            f"Oi, tudo bem? Vi a {nome} no Google e percebi que vocês têm boas "
            f"avaliações. Notei que uma página simples com cortes, serviços, horário, "
            f"localização e botão direto para agendar pelo WhatsApp poderia facilitar "
            f"muito para novos clientes. Quando o cliente precisa procurar informação "
            f"demais, ele pode acabar chamando outra barbearia. Eu montei uma prévia "
            f"visual de como uma página assim poderia ficar para uma barbearia. Posso te mandar?"
        )

    # Comida / Cardápio
    if any(n in nicho for n in NICHOS_COMIDA):
        return (
            f"Oi, tudo bem? Vi a {nome} no Google e percebi que vocês têm boas "
            f"avaliações. Notei que um cardápio digital simples, com produtos, fotos, "
            f"localização e botão direto para pedido no WhatsApp, poderia facilitar "
            f"para quem encontra vocês pelo Google ou Instagram. Quando o cliente "
            f"precisa procurar muito para ver o cardápio ou pedir, ele pode acabar "
            f"comprando em outro lugar. Eu montei uma prévia visual de como ficaria "
            f"um cardápio digital nesse estilo. Posso te mandar?"
        )

    # Dentista / Clínica
    if any(n in nicho for n in NICHOS_DENTISTA):
        return (
            f"Oi, tudo bem? Vi a {nome} no Google e percebi que vocês têm boas "
            f"avaliações. Notei que uma página profissional com tratamentos, localização, "
            f"fotos da clínica e botão direto para agendamento no WhatsApp poderia "
            f"passar ainda mais confiança para novos pacientes. Quando alguém procura "
            f"atendimento e não encontra as informações rápido, pode acabar escolhendo "
            f"outra clínica. Eu montei uma prévia visual de como uma página assim "
            f"poderia ficar. Posso te mandar?"
        )

    # Advocacia
    if any(n in nicho for n in NICHOS_ADVOCACIA):
        return (
            f"Olá, tudo bem? Vi o perfil da {nome} e percebi que uma página "
            f"institucional simples poderia ajudar a apresentar melhor as áreas de "
            f"atuação, localização e canal de contato pelo WhatsApp. Para serviços "
            f"jurídicos, uma presença online organizada transmite mais confiança para "
            f"quem está buscando orientação. Eu montei uma prévia visual de como uma "
            f"página profissional poderia ficar para um escritório de advocacia. "
            f"Posso te mandar?"
        )

    # Assistência técnica / Serviços
    if any(n in nicho for n in NICHOS_SERVICOS):
        return (
            f"Oi, tudo bem? Vi a {nome} no Google e percebi que vocês atendem um "
            f"serviço que muita gente procura com urgência. Uma página simples com "
            f"serviços, região atendida, localização e botão direto para orçamento "
            f"no WhatsApp pode facilitar muito para quem precisa chamar rápido. "
            f"Quando o cliente não encontra essas informações de forma clara, ele "
            f"pode chamar outro profissional. Eu montei uma prévia visual de como "
            f"ficaria. Posso te mandar?"
        )

    # Automotiva / Oficina
    if any(n in nicho for n in NICHOS_AUTOMOTIVA):
        return (
            f"Oi, tudo bem? Vi a {nome} no Google e percebi que vocês têm um serviço "
            f"com bastante apelo visual. Uma página simples com serviços, fotos de "
            f"antes e depois, localização e botão direto para orçamento no WhatsApp "
            f"poderia passar mais confiança para quem está procurando cuidado para "
            f"o carro. Eu montei uma prévia visual de como uma página assim poderia "
            f"ficar. Posso te mandar?"
        )

    # Igreja / Evento
    if any(n in nicho for n in NICHOS_IGREJA):
        return (
            f"Oi, tudo bem? Vi a {nome} e pensei em uma ideia simples: uma página "
            f"para divulgar eventos, cultos, campanhas ou programações com data, "
            f"local, informações e botão direto para WhatsApp ou inscrição. Isso "
            f"deixa tudo organizado em um link só para compartilhar com membros e "
            f"visitantes. Eu montei uma prévia visual de como uma página assim "
            f"poderia ficar. Posso te mandar?"
        )

    # Cursos / Reforço
    if any(n in nicho for n in NICHOS_CURSO):
        return (
            f"Oi, tudo bem? Vi a {nome} no Google e percebi que vocês trabalham com "
            f"ensino/atendimento educacional. Uma página simples com cursos, turmas, "
            f"localização, benefícios e botão direto para matrícula ou informações no "
            f"WhatsApp pode facilitar bastante para pais e alunos. Eu montei uma "
            f"prévia visual de como uma página assim poderia ficar. Posso te mandar?"
        )

    # Personalizados / Doces / Festas
    if any(n in nicho for n in NICHOS_PERSONALIZADOS):
        return (
            f"Oi, tudo bem? Vi a {nome} e percebi que vocês poderiam se beneficiar "
            f"de um catálogo digital simples, com produtos, fotos, prazos e botão "
            f"direto para orçamento pelo WhatsApp. Isso ajuda o cliente a ver as "
            f"opções sem ficar perguntando tudo no direct ou no WhatsApp. Eu montei "
            f"uma prévia visual de como um catálogo assim poderia ficar. Posso te mandar?"
        )

    # Mensagem geral
    return (
        f"Oi, tudo bem? Vi a {nome} no Google e percebi que vocês têm boas "
        f"avaliações. Notei que uma página simples com serviços, fotos, localização "
        f"e botão direto para WhatsApp poderia deixar tudo mais organizado para quem "
        f"encontra vocês pelo Google ou Instagram. Quando o cliente precisa procurar "
        f"muito, ele pode acabar desistindo ou chamando outro lugar. Eu montei uma "
        f"prévia visual de como uma página assim poderia ficar. Posso te mandar?"
    )


def gerar_link_whatsapp(lead):
    """Gera link wa.me com mensagem preenchida."""
    telefone = lead.get("whatsapp", "").strip()
    if not telefone:
        telefone = lead.get("telefone", "").strip()

    numeros = re.sub(r"\D", "", telefone)
    if not numeros:
        return ""

    if numeros.startswith("55") and len(numeros) >= 12:
        numero = numeros
    elif numeros.startswith("0"):
        numeros = numeros[1:]
        numero = "55" + numeros if len(numeros) in (10, 11) else ""
    elif len(numeros) == 11:
        numero = "55" + numeros
    elif len(numeros) == 10:
        numero = "55" + numeros
    else:
        return ""

    if not numero:
        return ""

    mensagem = lead.get("mensagem_whatsapp", "")
    if mensagem:
        return f"https://wa.me/{numero}?text={quote(mensagem)}"
    return f"https://wa.me/{numero}"


# ═══════════════════════════════════════════════════════════════════
# FILTROS
# ═══════════════════════════════════════════════════════════════════

def filtrar_leads(leads):
    """Filtra leads elegíveis para vender landing page."""
    filtrados = []
    for lead_raw in leads:
        lead = normalizar_lead(lead_raw)

        # Prioridade Alta
        if lead["prioridade"] != "Alta":
            continue

        # Score >= 70
        if lead["score"] < 70:
            continue

        # Deve ter telefone ou WhatsApp
        if not lead["telefone"] and not lead["whatsapp"]:
            continue

        # Sem site profissional
        if lead["site"] == "Sim" and lead["url_site"]:
            dominios_gratuitos = [
                "facebook.com", "instagram.com", "fb.me", "bit.ly",
                "wix.com", "wordpress.com", "blogspot.com",
            ]
            if not any(d in lead["url_site"].lower() for d in dominios_gratuitos):
                continue

        # Nicho bom para landing page
        nicho = lead["nicho"].lower()
        if not any(nb in nicho for nb in NICHOS_BONS):
            continue

        filtrados.append(lead)

    return filtrados


def ordenar_leads(leads):
    """Ordena por score, WhatsApp, avaliações, nota."""
    def chave(lead):
        tem_whatsapp = 1 if lead.get("whatsapp") or lead.get("link_whatsapp") else 0
        sem_site = 1 if lead.get("site") == "Não" else 0
        return (
            -lead.get("score", 0),
            -tem_whatsapp,
            -sem_site,
            -lead.get("qtd_avaliacoes", 0),
            -lead.get("nota_google", 0),
        )
    return sorted(leads, key=chave)


def classificar_nicho(lead):
    """Retorna categoria de filtro para o lead."""
    nicho = lead.get("nicho", "").lower()
    if any(n in nicho for n in NICHOS_ESTETICA):
        return "estetica"
    if any(n in nicho for n in NICHOS_BARBEARIA):
        return "barbearia"
    if any(n in nicho for n in NICHOS_COMIDA):
        return "comida"
    if any(n in nicho for n in NICHOS_DENTISTA):
        return "dentista"
    if any(n in nicho for n in NICHOS_ADVOCACIA):
        return "advocacia"
    if any(n in nicho for n in NICHOS_SERVICOS):
        return "servicos"
    if any(n in nicho for n in NICHOS_AUTOMOTIVA):
        return "automotiva"
    if any(n in nicho for n in NICHOS_IGREJA):
        return "igreja"
    if any(n in nicho for n in NICHOS_CURSO):
        return "cursos"
    if any(n in nicho for n in NICHOS_PERSONALIZADOS):
        return "personalizados"
    return "outros"


# ═══════════════════════════════════════════════════════════════════
# STATUS LEADS JSON
# ═══════════════════════════════════════════════════════════════════

def carregar_status():
    """Carrega status existente do JSON."""
    caminho = PAINEL_DIR / "status_leads.json"
    if caminho.exists():
        try:
            with open(caminho, "r", encoding="utf-8") as f:
                return json.load(f)
        except (json.JSONDecodeError, IOError):
            return {}
    return {}


def carregar_historico():
    """Carrega historico acumulado de todos os leads ja abordados."""
    caminho = PAINEL_DIR / "historico_leads.json"
    if caminho.exists():
        try:
            with open(caminho, "r", encoding="utf-8") as f:
                return json.load(f)
        except (json.JSONDecodeError, IOError):
            return {}
    return {}


def salvar_historico(historico):
    """Salva historico acumulado de leads."""
    caminho = PAINEL_DIR / "historico_leads.json"
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(historico, f, ensure_ascii=False, indent=2)
    return caminho


def mesclar_status(leads, status_existente):
    """Mescla dados dos leads com status salvo."""
    for lead in leads:
        lid = lead["lead_id"]
        if lid in status_existente:
            lead["status"] = status_existente[lid].get("status", "novo")
            lead["data_abordagem"] = status_existente[lid].get("data_abordagem", "")
            lead["data_followup"] = status_existente[lid].get("data_followup", "")
            lead["observacoes"] = status_existente[lid].get("observacoes", "")
    return leads


def salvar_status(leads):
    """Salva arquivo de controle de status."""
    status = {}
    for lead in leads:
        status[lead["lead_id"]] = {
            "lead_id": lead["lead_id"],
            "nome": lead["nome"],
            "status": lead.get("status", "novo"),
            "data_abordagem": lead.get("data_abordagem", ""),
            "data_followup": lead.get("data_followup", ""),
            "observacoes": lead.get("observacoes", ""),
        }

    caminho = PAINEL_DIR / "status_leads.json"
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(status, f, ensure_ascii=False, indent=2)

    return caminho


# ═══════════════════════════════════════════════════════════════════
# GERAÇÃO DO HTML
# ═══════════════════════════════════════════════════════════════════

def gerar_html(leads, data_geracao, metricas=None):
    """Gera o HTML completo do painel."""
    leads_json = json.dumps(leads, ensure_ascii=False)

    hoje_str = data_geracao.strftime("%d/%m/%Y")

    # Metricas acumuladas
    if metricas is None:
        metricas = {
            "total_abordados": len(leads),
            "mensagens_enviadas": 0,
            "responderam": 0,
            "interessados": 0,
            "propostas": 0,
            "fechados": 0,
        }

    metricas_json = json.dumps(metricas, ensure_ascii=False)

    html = HTML_TEMPLATE.replace("__LEADS_DATA__", leads_json, 1)
    html = html.replace("__DATA_GERACAO__", hoje_str, 1)
    html = html.replace("__TOTAL_LEADS__", str(len(leads)), 1)
    html = html.replace("__METRICAS_ACUMULADAS__", metricas_json, 1)

    return html


HTML_TEMPLATE = r"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Painel de Prospecção - __DATA_GERACAO__</title>
<style>
*{margin:0;padding:0;box-sizing:border-box}
:root{
  --bg-deep:#08090a;--bg-card:#111214;--bg-elevated:#18191c;--bg-hover:#1e2023;
  --accent:#6d5dfc;--accent-glow:rgba(109,93,252,.25);--accent2:#c084fc;
  --green:#34d399;--green-dim:rgba(52,211,153,.12);
  --blue:#60a5fa;--blue-dim:rgba(96,165,250,.12);
  --yellow:#fbbf24;--yellow-dim:rgba(251,191,36,.12);
  --red:#f87171;--red-dim:rgba(248,113,113,.12);
  --orange:#fb923c;--orange-dim:rgba(251,146,60,.12);
  --cyan:#22d3ee;--cyan-dim:rgba(34,211,238,.12);
  --purple:#a78bfa;--purple-dim:rgba(167,139,250,.12);
  --pink:#f472b6;--pink-dim:rgba(244,114,182,.12);
  --text:#f1f1f1;--text2:#a0a0a8;--muted:#6b6b76;
  --border:rgba(255,255,255,.06);--border-active:rgba(109,93,252,.4);
  --radius:12px;
}
html{scroll-behavior:smooth}
body{font-family:'Segoe UI',-apple-system,BlinkMacSystemFont,sans-serif;
  background:var(--bg-deep);color:var(--text);min-height:100vh;line-height:1.6}
.bg-mesh{position:fixed;inset:0;pointer-events:none;z-index:0;opacity:.35;
  background-image:
    radial-gradient(ellipse at 10% 10%,rgba(109,93,252,.08) 0%,transparent 50%),
    radial-gradient(ellipse at 90% 80%,rgba(34,211,238,.06) 0%,transparent 50%),
    radial-gradient(ellipse at 50% 50%,rgba(52,211,153,.04) 0%,transparent 70%)}
.container{position:relative;z-index:1;max-width:1500px;margin:0 auto;padding:0 20px 60px}

/* Header */
header{padding:32px 0 24px;border-bottom:1px solid var(--border);margin-bottom:24px}
.header-top{display:flex;justify-content:space-between;align-items:center;gap:20px;flex-wrap:wrap}
.header-title h1{font-size:1.8rem;font-weight:700;letter-spacing:-.02em;
  background:linear-gradient(135deg,var(--accent),var(--accent2),var(--cyan));
  -webkit-background-clip:text;-webkit-text-fill-color:transparent;background-clip:text}
.header-title p{color:var(--text2);font-size:.9rem;margin-top:4px}
.header-date{color:var(--muted);font-size:.8rem;margin-top:2px}

/* Summary cards */
.summary-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:20px 0 24px}
.summary-card{background:var(--bg-card);border:1px solid var(--border);border-radius:var(--radius);
  padding:16px;text-align:center;transition:all .2s;cursor:default}
.summary-card:hover{border-color:var(--border-active);transform:translateY(-2px);
  box-shadow:0 8px 24px rgba(0,0,0,.3)}
.summary-value{font-size:2rem;font-weight:700;line-height:1;margin-bottom:4px}
.summary-label{font-size:.7rem;color:var(--muted);text-transform:uppercase;letter-spacing:.08em}
.summary-card.leads .summary-value{color:var(--accent)}
.summary-card.enviadas .summary-value{color:var(--cyan)}
.summary-card.responderam .summary-value{color:var(--green)}
.summary-card.interessados .summary-value{color:var(--yellow)}
.summary-card.propostas .summary-value{color:var(--orange)}
.summary-card.fechados .summary-value{color:var(--green)}
.summary-card.taxa .summary-value{color:var(--purple)}

/* Filters */
.filters{display:flex;flex-wrap:wrap;gap:8px;margin-bottom:20px;padding:16px;
  background:var(--bg-card);border-radius:var(--radius);border:1px solid var(--border)}
.filter-btn{background:var(--bg-elevated);border:1px solid var(--border);border-radius:8px;
  padding:7px 16px;color:var(--text2);font-size:.82rem;cursor:pointer;transition:all .2s;
  font-family:inherit;white-space:nowrap}
.filter-btn:hover{border-color:var(--accent);color:var(--text)}
.filter-btn.active{background:var(--accent);border-color:var(--accent);color:#fff;font-weight:600}
.search-box{flex:1;min-width:200px}
.search-input{width:100%;background:var(--bg-elevated);border:1px solid var(--border);
  border-radius:8px;padding:8px 14px;color:var(--text);font-size:.85rem;font-family:inherit;
  transition:all .2s}
.search-input:focus{outline:none;border-color:var(--accent);box-shadow:0 0 0 3px var(--accent-glow)}
.search-input::placeholder{color:var(--muted)}

/* Leads count */
.results-info{display:flex;justify-content:space-between;align-items:center;margin-bottom:16px;
  font-size:.85rem;color:var(--text2)}
.results-info strong{color:var(--accent);font-weight:600}

/* Lead cards */
.leads-grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(420px,1fr));gap:16px;margin-bottom:40px}

.lead-card{background:var(--bg-card);border:1px solid var(--border);border-radius:var(--radius);
  overflow:hidden;transition:all .3s cubic-bezier(.4,0,.2,1);position:relative;
  opacity:0;animation:fadeUp .4s ease forwards}
.lead-card:hover{transform:translateY(-3px);border-color:var(--border-active);
  box-shadow:0 12px 32px rgba(0,0,0,.3),0 0 40px var(--accent-glow)}

.lead-card.alta::before{content:'';position:absolute;top:0;left:0;right:0;height:3px;
  background:linear-gradient(90deg,var(--accent),var(--accent2))}
.lead-card.media::before{content:'';position:absolute;top:0;left:0;right:0;height:3px;
  background:linear-gradient(90deg,var(--yellow),var(--orange))}

.card-header{padding:14px 16px 8px;display:flex;justify-content:space-between;align-items:flex-start;gap:8px}
.card-badges{display:flex;gap:5px;flex-wrap:wrap}
.badge{font-size:.6rem;text-transform:uppercase;letter-spacing:.06em;padding:2px 8px;
  border-radius:100px;font-weight:600;white-space:nowrap}
.badge-nicho{background:var(--purple-dim);color:var(--purple)}
.badge-cidade{background:var(--cyan-dim);color:var(--cyan)}
.badge-sem-site{background:var(--green-dim);color:var(--green)}
.badge-whatsapp{background:var(--blue-dim);color:var(--blue)}
.badge-alta{background:rgba(109,93,252,.15);color:var(--accent)}
.badge-score{font-size:.6rem;padding:2px 8px;border-radius:100px;font-weight:700}
.badge-score.high{background:var(--green-dim);color:var(--green)}
.badge-score.mid{background:var(--yellow-dim);color:var(--yellow)}

.card-body{padding:0 16px 12px}
.card-name{font-size:1.05rem;font-weight:600;line-height:1.3;margin-bottom:6px;
  display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden}
.card-meta{display:flex;flex-direction:column;gap:5px;margin-bottom:10px}
.meta-row{display:flex;align-items:center;gap:6px;font-size:.82rem;color:var(--text2)}
.meta-icon{width:14px;height:14px;flex-shrink:0;opacity:.5}
.meta-icon svg{width:100%;height:100%}

.card-info-grid{display:grid;grid-template-columns:1fr 1fr;gap:4px 12px;margin-bottom:10px;font-size:.78rem;color:var(--text2)}
.card-info-grid .info-label{color:var(--muted);font-size:.7rem;text-transform:uppercase;letter-spacing:.04em}
.card-info-grid .info-value{color:var(--text2)}

/* Message box */
.msg-box{background:var(--bg-elevated);border:1px solid var(--border);border-radius:8px;
  padding:10px 12px;margin-bottom:10px;position:relative}
.msg-text{font-size:.8rem;color:var(--text2);line-height:1.5;max-height:60px;overflow:hidden;
  transition:max-height .3s}
.msg-text.expanded{max-height:500px}
.msg-toggle{font-size:.7rem;color:var(--accent);cursor:pointer;margin-top:4px;display:inline-block}
.msg-toggle:hover{text-decoration:underline}
.copy-btn{position:absolute;top:8px;right:8px;background:var(--accent);border:none;border-radius:6px;
  padding:4px 10px;color:#fff;font-size:.7rem;cursor:pointer;transition:all .2s;font-family:inherit}
.copy-btn:hover{opacity:.85;transform:scale(1.02)}
.copy-btn.copied{background:var(--green)}

/* Actions */
.card-actions{display:flex;flex-wrap:wrap;gap:6px;padding:0 16px 14px}
.action-btn{background:var(--bg-elevated);border:1px solid var(--border);border-radius:8px;
  padding:5px 10px;color:var(--text2);font-size:.72rem;cursor:pointer;transition:all .2s;
  font-family:inherit;white-space:nowrap}
.action-btn:hover{border-color:var(--accent);color:var(--text)}
.action-btn.whatsapp{background:rgba(37,211,102,.1);border-color:rgba(37,211,102,.3);color:#25d366}
.action-btn.whatsapp:hover{background:rgba(37,211,102,.2);border-color:rgba(37,211,102,.5)}
.action-btn.enviada{background:var(--cyan-dim);border-color:rgba(96,165,250,.3);color:var(--blue)}
.action-btn.respondeu{background:var(--green-dim);border-color:rgba(52,211,153,.3);color:var(--green)}
.action-btn.video{background:var(--purple-dim);border-color:rgba(167,139,250,.3);color:var(--purple)}
.action-btn.proposta{background:var(--orange-dim);border-color:rgba(251,146,60,.3);color:var(--orange)}
.action-btn.fechado{background:var(--green-dim);border-color:rgba(52,211,153,.3);color:var(--green)}
.action-btn.perdido{background:var(--red-dim);border-color:rgba(248,113,113,.3);color:var(--red)}
.action-btn.followup{background:var(--yellow-dim);border-color:rgba(251,191,36,.3);color:var(--yellow)}

/* Status bar */
.status-bar{padding:8px 16px 12px;display:flex;align-items:center;gap:8px}
.status-badge{font-size:.7rem;padding:3px 10px;border-radius:100px;font-weight:600;text-transform:uppercase}
.status-badge.novo{background:var(--bg-elevated);color:var(--muted);border:1px solid var(--border)}
.status-badge.mensagem-enviada{background:var(--cyan-dim);color:var(--cyan)}
.status-badge.respondeu{background:var(--green-dim);color:var(--green)}
.status-badge.video-enviado{background:var(--purple-dim);color:var(--purple)}
.status-badge.interessado{background:var(--yellow-dim);color:var(--yellow)}
.status-badge.proposta-enviada{background:var(--orange-dim);color:var(--orange)}
.status-badge.fechado{background:rgba(52,211,153,.2);color:var(--green);border:1px solid rgba(52,211,153,.3)}
.status-badge.perdido{background:var(--red-dim);color:var(--red)}
.status-badge.follow-up{background:var(--yellow-dim);color:var(--yellow)}
.status-badge.abordar-hoje{background:var(--accent);color:#fff;background:rgba(109,93,252,.15);color:var(--accent)}
.status-badge.nao-abordar{background:rgba(107,107,118,.1);color:var(--muted)}

/* Observations */
.obs-section{padding:0 16px 14px}
.obs-input{width:100%;background:var(--bg-elevated);border:1px solid var(--border);border-radius:6px;
  padding:6px 10px;color:var(--text);font-size:.78rem;font-family:inherit;resize:vertical;min-height:32px}
.obs-input:focus{outline:none;border-color:var(--accent)}

/* Empty state */
.empty-state{text-align:center;padding:60px 20px;color:var(--muted);grid-column:1/-1}
.empty-state h3{font-size:1.2rem;color:var(--text2);margin-bottom:6px}

/* Animations */
@keyframes fadeUp{from{opacity:0;transform:translateY(12px)}to{opacity:1;transform:translateY(0)}}

/* Toast */
.toast{position:fixed;bottom:24px;right:24px;background:var(--green);color:#fff;padding:10px 20px;
  border-radius:8px;font-size:.85rem;font-weight:600;z-index:1000;opacity:0;transform:translateY(10px);
  transition:all .3s;pointer-events:none}
.toast.show{opacity:1;transform:translateY(0)}

/* Footer */
footer{padding:24px 0;border-top:1px solid var(--border);text-align:center;color:var(--muted);font-size:.8rem}

/* Responsive */
@media(max-width:900px){
  .leads-grid{grid-template-columns:1fr}
  .summary-grid{grid-template-columns:repeat(auto-fit,minmax(120px,1fr))}
  .header-top{flex-direction:column;align-items:flex-start}
}
@media(max-width:500px){
  .filters{padding:10px;gap:6px}
  .filter-btn{padding:6px 10px;font-size:.75rem}
  .card-actions{gap:4px}
  .action-btn{padding:4px 7px;font-size:.68rem}
}
</style>
</head>
<body>
<div class="bg-mesh"></div>
<div class="container">

<header>
  <div class="header-top">
    <div class="header-title">
      <h1>Painel de Prospecção</h1>
      <p>Leads diários para vender landing pages e mini sites</p>
      <div class="header-date">Gerado em: __DATA_GERACAO__</div>
    </div>
  </div>
</header>

<section class="summary-grid" id="summaryGrid">
  <div class="summary-card leads"><div class="summary-value" id="sLeads">0</div><div class="summary-label">Leads para hoje</div></div>
  <div class="summary-card enviadas"><div class="summary-value" id="sEnviadas">0</div><div class="summary-label">Mensagens enviadas</div></div>
  <div class="summary-card responderam"><div class="summary-value" id="sResponderam">0</div><div class="summary-label">Responderam</div></div>
  <div class="summary-card interessados"><div class="summary-value" id="sInteressados">0</div><div class="summary-label">Interessados</div></div>
  <div class="summary-card propostas"><div class="summary-value" id="sPropostas">0</div><div class="summary-label">Propostas enviadas</div></div>
  <div class="summary-card fechados"><div class="summary-value" id="sFechados">0</div><div class="summary-label">Fechados</div></div>
  <div class="summary-card taxa"><div class="summary-value" id="sTaxa">0%</div><div class="summary-label">Taxa de resposta</div></div>
</section>

<section class="filters">
  <button class="filter-btn active" data-filter="todos">Todos</button>
  <button class="filter-btn" data-filter="alta">Alta prioridade</button>
  <button class="filter-btn" data-filter="sem-site">Sem site</button>
  <button class="filter-btn" data-filter="com-whatsapp">Com WhatsApp</button>
  <button class="filter-btn" data-filter="estetica">Estética</button>
  <button class="filter-btn" data-filter="barbearia">Barbearia</button>
  <button class="filter-btn" data-filter="comida">Comida/Cardápio</button>
  <button class="filter-btn" data-filter="dentista">Dentista/Clínica</button>
  <button class="filter-btn" data-filter="advocacia">Advocacia</button>
  <button class="filter-btn" data-filter="servicos">Serviços técnicos</button>
  <button class="filter-btn" data-filter="automotiva">Automotiva</button>
  <button class="filter-btn" data-filter="igreja">Igreja/Evento</button>
  <button class="filter-btn" data-filter="cursos">Cursos</button>
  <button class="filter-btn" data-filter="personalizados">Personalizados</button>
  <button class="filter-btn" data-filter="respondidos">Respondidos</button>
  <button class="filter-btn" data-filter="followup">Follow-up</button>
  <button class="filter-btn" data-filter="fechados">Fechados</button>
  <button class="filter-btn" data-filter="perdidos">Perdidos</button>
  <div class="search-box">
    <input type="text" class="search-input" id="searchInput" placeholder="Buscar por nome, nicho ou cidade...">
  </div>
</section>

<div class="results-info">
  <span>Mostrando <strong id="visibleCount">0</strong> de <strong id="totalFiltered">0</strong> leads</span>
</div>

<main class="leads-grid" id="leadsGrid"></main>

<footer>
  <p>Painel de Prospecção — Atualize o status de cada lead após cada contato</p>
  <p style="margin-top:4px">Este painel NÃO envia mensagens automaticamente. Use os botões para copiar e abrir WhatsApp manualmente.</p>
</footer>

</div>

<div class="toast" id="toast"></div>

<script>
const LEADS = __LEADS_DATA__;
const METRICAS_ACUMULADAS = __METRICAS_ACUMULADAS__;
let currentFilter = 'todos';
let searchTerm = '';

// ── Status persistence (SQLite via API) ────────────────────────
let statusCache = null;

async function loadStatus() {
  if (statusCache) return statusCache;
  try {
    const res = await fetch('/api/leads');
    const data = await res.json();
    const map = {};
    (data.leads || []).forEach(l => {
      map[l.lead_id] = {
        status: l.status || 'novo',
        data_abordagem: l.data_abordagem || '',
        data_followup: l.data_followup || '',
        observacoes: l.observacoes || ''
      };
    });
    statusCache = map;
    return map;
  } catch(e) {
    console.warn('Erro ao carregar status do servidor, usando localStorage:', e);
    try {
      const raw = localStorage.getItem('prospeccao_status_v2');
      return raw ? JSON.parse(raw) : {};
    } catch { return {}; }
  }
}

function saveStatusLocal(status) {
  localStorage.setItem('prospeccao_status_v2', JSON.stringify(status));
}

async function setStatus(leadId, field, value) {
  // Atualiza local imediatamente para UI rapida
  const s = await loadStatus();
  if (!s[leadId]) s[leadId] = { status: 'novo', data_abordagem: '', data_followup: '', observacoes: '' };
  s[leadId][field] = value;
  statusCache = s;
  saveStatusLocal(s);
  render();

  // Salva no servidor (SQLite)
  try {
    const res = await fetch('/api/status-single', {
      method: 'POST',
      headers: {'Content-Type': 'application/json'},
      body: JSON.stringify({ lead_id: leadId, ...s[leadId] })
    });
    const data = await res.json();
    if (data.ok && data.metricas) {
      updateSummaryFromMetricas(data.metricas);
    }
  } catch(e) {
    console.warn('Erro ao salvar no servidor:', e);
  }
}

async function updateObs(leadId, value) {
  return setStatus(leadId, 'observacoes', value);
}

function updateSummaryFromMetricas(m) {
  animNum('sLeads', m.total_abordados);
  animNum('sEnviadas', m.mensagens_enviadas);
  animNum('sResponderam', m.responderam);
  animNum('sInteressados', m.interessados);
  animNum('sPropostas', m.propostas);
  animNum('sFechados', m.fechados);
  const taxa = m.mensagens_enviadas > 0 ? Math.round((m.responderam / m.mensagens_enviadas) * 100) : 0;
  document.getElementById('sTaxa').textContent = taxa + '%';
}

// ── Filtering ────────────────────────────────────────────────
async function getFiltered() {
  const status = await loadStatus();
  let filtered = LEADS.filter(l => {
    const s = status[l.lead_id] || {};
    const st = s.status || 'novo';
    const q = searchTerm.toLowerCase();

    // Search
    if (q && !(l.nome.toLowerCase().includes(q) || l.nicho.toLowerCase().includes(q) || l.cidade.toLowerCase().includes(q) || (l.bairro||'').toLowerCase().includes(q))) return false;

    // Status filters
    if (currentFilter === 'respondidos') return ['respondeu','interessado'].includes(st);
    if (currentFilter === 'followup') return st === 'follow-up';
    if (currentFilter === 'fechados') return st === 'fechado';
    if (currentFilter === 'perdidos') return st === 'perdido';

    // Niche filters
    if (currentFilter === 'estetica') return l.categoria === 'estetica';
    if (currentFilter === 'barbearia') return l.categoria === 'barbearia';
    if (currentFilter === 'comida') return l.categoria === 'comida';
    if (currentFilter === 'dentista') return l.categoria === 'dentista';
    if (currentFilter === 'advocacia') return l.categoria === 'advocacia';
    if (currentFilter === 'servicos') return l.categoria === 'servicos';
    if (currentFilter === 'automotiva') return l.categoria === 'automotiva';
    if (currentFilter === 'igreja') return l.categoria === 'igreja';
    if (currentFilter === 'cursos') return l.categoria === 'cursos';
    if (currentFilter === 'personalizados') return l.categoria === 'personalizados';

    // Attribute filters
    if (currentFilter === 'alta') return l.prioridade === 'Alta';
    if (currentFilter === 'sem-site') return l.site === 'Não';
    if (currentFilter === 'com-whatsapp') return !!(l.whatsapp || l.link_whatsapp);

    // "Todos" mostra apenas leads que ainda nao foram abordados
    const naoAbordado = ['novo', 'abordar hoje'].includes(st);
    return naoAbordado;
  });
  return filtered;
}

// ── Summary ───────────────────────────────────────────────────
async function updateSummary() {
  const status = await loadStatus();
  let enviadas = 0, responderam = 0, interessados = 0, propostas = 0, fechados = 0;
  LEADS.forEach(l => {
    const s = (status[l.lead_id] || {}).status || 'novo';
    if (['mensagem enviada','video enviado','respondeu','interessado','proposta enviada','fechado','follow-up'].includes(s)) enviadas++;
    if (['respondeu','interessado','proposta enviada','fechado'].includes(s)) responderam++;
    if (['interessado','proposta enviada','fechado'].includes(s)) interessados++;
    if (['proposta enviada','fechado'].includes(s)) propostas++;
    if (s === 'fechado') fechados++;
  });
  const taxa = enviadas > 0 ? Math.round((responderam / enviadas) * 100) : 0;
  animNum('sLeads', LEADS.length);
  animNum('sEnviadas', enviadas);
  animNum('sResponderam', responderam);
  animNum('sInteressados', interessados);
  animNum('sPropostas', propostas);
  animNum('sFechados', fechados);
  document.getElementById('sTaxa').textContent = taxa + '%';
}

function animNum(id, target) {
  const el = document.getElementById(id);
  const start = performance.now();
  (function tick(now) {
    const p = Math.min((now - start) / 600, 1);
    el.textContent = Math.floor(target * (1 - Math.pow(1 - p, 3)));
    if (p < 1) requestAnimationFrame(tick);
  })(start);
}

// ── Icons ─────────────────────────────────────────────────────
const IC = {
  phone: '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><path d="M22 16.92v3a2 2 0 01-2.18 2 19.79 19.79 0 01-8.63-3.07 19.5 19.5 0 01-6-6 19.79 19.79 0 01-3.07-8.67A2 2 0 014.11 2h3a2 2 0 012 1.72c.127.96.361 1.903.7 2.81a2 2 0 01-.45 2.11L8.09 9.91a16 16 0 006 6l1.27-1.27a2 2 0 012.11-.45c.907.339 1.85.573 2.81.7A2 2 0 0122 16.92z"/></svg>',
  pin: '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><path d="M21 10c0 7-9 13-9 13s-9-6-9-13a9 9 0 0118 0z"/><circle cx="12" cy="10" r="3"/></svg>',
  star: '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><path d="M12 2l3.09 6.26L22 9.27l-5 4.87 1.18 6.88L12 17.77l-6.18 3.25L7 14.14 2 9.27l6.91-1.01L12 2z"/></svg>',
  wa: '<svg viewBox="0 0 24 24" fill="currentColor"><path d="M17.472 14.382c-.297-.149-1.758-.867-2.03-.967-.273-.099-.471-.148-.67.15-.197.297-.767.966-.94 1.164-.173.199-.347.223-.644.075-.297-.15-1.255-.463-2.39-1.475-.883-.788-1.48-1.761-1.653-2.059-.173-.297-.018-.458.13-.606.134-.133.298-.347.446-.52.149-.174.198-.298.298-.497.099-.198.05-.371-.025-.52-.075-.149-.669-1.612-.916-2.207-.242-.579-.487-.5-.669-.51-.173-.008-.371-.01-.57-.01-.198 0-.52.074-.792.372-.272.297-1.04 1.016-1.04 2.479 0 1.462 1.065 2.875 1.213 3.074.149.198 2.096 3.2 5.077 4.487.709.306 1.262.489 1.694.625.712.227 1.36.195 1.871.118.571-.085 1.758-.719 2.006-1.413.248-.694.248-1.289.173-1.413-.074-.124-.272-.198-.57-.347m-5.421 7.403h-.004a9.87 9.87 0 01-5.031-1.378l-.361-.214-3.741.981.998-3.648-.235-.374a9.86 9.86 0 01-1.51-5.26c.001-5.42 4.432-9.849 9.855-9.849 2.633 0 5.117 1.027 6.988 2.891a9.825 9.825 0 012.89 6.994c-.003 5.42-4.433 9.849-9.855 9.849m8.398-18.276A13.892 13.892 0 0012.047 0C5.395 0 .028 5.362.025 12.017c-.001 2.118.553 4.182 1.61 6.003L0 24l6.095-1.597a13.867 13.867 0 006.426 1.581h.005c6.548 0 11.917-5.362 11.92-11.917a11.834 11.834 0 00-3.478-8.392"/></svg>',
};

// ── Render ────────────────────────────────────────────────────
function esc(s) { const d = document.createElement('div'); d.textContent = s; return d.innerHTML; }

function statusClass(st) {
  const map = {
    'novo':'novo','abordar hoje':'abordar-hoje','mensagem enviada':'mensagem-enviada',
    'respondeu':'respondeu','video enviado':'video-enviado','interessado':'interessado',
    'proposta enviada':'proposta-enviada','fechado':'fechado','perdido':'perdido',
    'follow-up':'follow-up','não abordar':'nao-abordar'
  };
  return map[st] || 'novo';
}

async function render() {
  const filtered = await getFiltered();
  const status = loadStatus();

  document.getElementById('visibleCount').textContent = filtered.length;
  document.getElementById('totalFiltered').textContent = LEADS.length;

  const grid = document.getElementById('leadsGrid');
  if (!filtered.length) {
    grid.innerHTML = '<div class="empty-state"><h3>Nenhum lead encontrado</h3><p>Ajuste os filtros para ver resultados.</p></div>';
    updateSummary();
    return;
  }

  grid.innerHTML = filtered.map((l, i) => {
    const s = status[l.lead_id] || {};
    const st = s.status || 'novo';
    const obs = s.observacoes || '';
    const sc = statusClass(st);
    const priClass = l.prioridade === 'Alta' ? 'alta' : 'media';

    // Badges
    let badges = '';
    badges += '<span class="badge badge-nicho">' + esc(l.nicho) + '</span>';
    badges += '<span class="badge badge-cidade">' + esc(l.cidade) + '</span>';
    if (l.site === 'Não') badges += '<span class="badge badge-sem-site">Sem site</span>';
    if (l.whatsapp || l.link_whatsapp) badges += '<span class="badge badge-whatsapp">WhatsApp</span>';
    if (l.prioridade === 'Alta') badges += '<span class="badge badge-alta">Alta</span>';

    // Score badge
    const scoreClass = l.score >= 85 ? 'high' : 'mid';
    badges += '<span class="badge badge-score ' + scoreClass + '">' + l.score + '</span>';

    // Meta rows
    let meta = '';
    if (l.telefone || l.whatsapp) {
      const fone = l.whatsapp || l.telefone;
      meta += '<div class="meta-row"><span class="meta-icon">' + IC.phone + '</span>' + esc(fone) + '</div>';
    }
    if (l.bairro) {
      meta += '<div class="meta-row"><span class="meta-icon">' + IC.pin + '</span>' + esc(l.bairro) + ' — ' + esc(l.cidade) + '</div>';
    }

    // Info grid
    let info = '';
    if (l.nota_google) info += '<div><div class="info-label">Nota Google</div><div class="info-value">' + l.nota_google.toFixed(1) + ' (' + l.qtd_avaliacoes + ')</div></div>';
    if (l.oferta_sugerida) info += '<div><div class="info-label">Oferta</div><div class="info-value">' + esc(l.oferta_sugerida) + '</div></div>';
    if (l.tipo_de_material) info += '<div><div class="info-label">Material</div><div class="info-value">' + esc(l.tipo_de_material) + '</div></div>';
    if (l.motivo_prioridade) info += '<div><div class="info-label">Motivo</div><div class="info-value">' + esc(l.motivo_prioridade) + '</div></div>';

    // WhatsApp link
    const waLink = l.link_whatsapp || '#';

    return '<article class="lead-card ' + priClass + '" style="animation-delay:' + (i * 30) + 'ms">' +
      '<div class="card-header"><div class="card-badges">' + badges + '</div></div>' +
      '<div class="card-body">' +
        '<div class="card-name">' + esc(l.nome) + '</div>' +
        '<div class="card-meta">' + meta + '</div>' +
        (info ? '<div class="card-info-grid">' + info + '</div>' : '') +
      '</div>' +
      (l.mensagem_whatsapp ? '<div class="msg-box"><button class="copy-btn" onclick="copyMsg(\'' + l.lead_id + '\',this)">Copiar</button><div class="msg-text" id="msg-' + l.lead_id + '">' + esc(l.mensagem_whatsapp) + '</div><span class="msg-toggle" onclick="toggleMsg(\'' + l.lead_id + '\',this)">Ver mais</span></div>' : '') +
      '<div class="card-actions">' +
        (waLink !== '#' ? '<a href="' + esc(waLink) + '" target="_blank" class="action-btn whatsapp">Abrir WhatsApp</a>' : '') +
        '<button class="action-btn enviada" onclick="setStatus(\'' + l.lead_id + '\',\'status\',\'mensagem enviada\')">Mensagem enviada</button>' +
        '<button class="action-btn respondeu" onclick="setStatus(\'' + l.lead_id + '\',\'status\',\'respondeu\')">Respondeu</button>' +
        '<button class="action-btn video" onclick="setStatus(\'' + l.lead_id + '\',\'status\',\'video enviado\')">Vídeo enviado</button>' +
        '<button class="action-btn proposta" onclick="setStatus(\'' + l.lead_id + '\',\'status\',\'proposta enviada\')">Proposta</button>' +
        '<button class="action-btn fechado" onclick="setStatus(\'' + l.lead_id + '\',\'status\',\'fechado\')">Fechado</button>' +
        '<button class="action-btn perdido" onclick="setStatus(\'' + l.lead_id + '\',\'status\',\'perdido\')">Perdido</button>' +
        '<button class="action-btn followup" onclick="setStatus(\'' + l.lead_id + '\',\'status\',\'follow-up\')">Follow-up</button>' +
      '</div>' +
      '<div class="status-bar"><span class="status-badge ' + sc + '">' + esc(st) + '</span></div>' +
      '<div class="obs-section"><textarea class="obs-input" placeholder="Observações..." onblur="setStatus(\'' + l.lead_id + '\',\'observacoes\',this.value)">' + esc(obs) + '</textarea></div>'
    '</article>';
  }).join('');

  updateSummary();
}

function copyMsg(leadId, btn) {
  const lead = LEADS.find(l => l.lead_id === leadId);
  if (!lead || !lead.mensagem_whatsapp) return;
  navigator.clipboard.writeText(lead.mensagem_whatsapp).then(() => {
    btn.textContent = 'Copiado!';
    btn.classList.add('copied');
    showToast('Mensagem copiada!');
    setTimeout(() => { btn.textContent = 'Copiar'; btn.classList.remove('copied'); }, 2000);
  }).catch(() => {
    // Fallback
    const ta = document.createElement('textarea');
    ta.value = lead.mensagem_whatsapp;
    document.body.appendChild(ta);
    ta.select();
    document.execCommand('copy');
    document.body.removeChild(ta);
    btn.textContent = 'Copiado!';
    btn.classList.add('copied');
    showToast('Mensagem copiada!');
    setTimeout(() => { btn.textContent = 'Copiar'; btn.classList.remove('copied'); }, 2000);
  });
}

function toggleMsg(leadId, toggle) {
  const el = document.getElementById('msg-' + leadId);
  if (el.classList.contains('expanded')) {
    el.classList.remove('expanded');
    toggle.textContent = 'Ver mais';
  } else {
    el.classList.add('expanded');
    toggle.textContent = 'Ver menos';
  }
}

function showToast(msg) {
  const t = document.getElementById('toast');
  t.textContent = msg;
  t.classList.add('show');
  setTimeout(() => t.classList.remove('show'), 2500);
}

// ── Events ────────────────────────────────────────────────────
document.querySelectorAll('.filter-btn').forEach(btn => {
  btn.addEventListener('click', () => {
    document.querySelectorAll('.filter-btn').forEach(b => b.classList.remove('active'));
    btn.classList.add('active');
    currentFilter = btn.dataset.filter;
    render();
  });
});

document.getElementById('searchInput').addEventListener('input', (e) => {
  searchTerm = e.target.value;
  render();
});

// ── Init ──────────────────────────────────────────────────────
loadStatus().then(() => render());
</script>
</body>
</html>"""


# ═══════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Painel de Prospecção - Gerador de Dashboard HTML")
    parser.add_argument("--limite", type=int, default=DAILY_LIMIT, help=f"Limite de leads por dia (padrão: {DAILY_LIMIT})")
    parser.add_argument("--novos", action="store_true", help="Excluir leads ja abordados e gerar painel com leads novos")
    args = parser.parse_args()

    if sys.platform == "win32":
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass

    print("=" * 60)
    print("  PAINEL DE PROSPECÇÃO - GERADOR DE DASHBOARD")
    print("=" * 60)

    # Encontra planilha
    caminho, tipo = encontrar_planilha()
    if not caminho:
        print("\n[!] Nenhuma planilha encontrada.")
        print("    Execute prospectar_leads.py ou campanha_diaria.py primeiro.")
        sys.exit(1)

    print(f"\n  Lendo: {caminho}")

    # Lê leads
    if tipo == "campanha":
        leads_raw = ler_todas_abas(caminho)
    else:
        leads_raw = ler_excel(caminho)

    print(f"  Leads carregados: {len(leads_raw)}")

    # Filtra
    leads_filtrados = filtrar_leads(leads_raw)
    print(f"  Leads elegíveis (Alta + score≥70 + contato + sem site): {len(leads_filtrados)}")

    if not leads_filtrados:
        print("\n[!] Nenhum lead elegível encontrado.")
        sys.exit(0)

    # Inicializar banco de dados
    sys.path.insert(0, str(PAINEL_DIR))
    import db
    db.init_db()
    conn = db.get_conn()

    # Determinar rodada atual
    rodada_atual = db.get_rodada_atual(conn) + 1
    if not args.novos:
        rodada_atual = max(1, rodada_atual - 1)  # Se nao for --novos, usar mesma rodada

    # Excluir leads ja abordados se --novos
    if args.novos:
        leads_existentes = conn.execute("SELECT lead_id FROM leads").fetchall()
        abordados = set(r["lead_id"] for r in leads_existentes)
        leads_filtrados = [l for l in leads_filtrados if l["lead_id"] not in abordados]
        print(f"  Leads ja abordados (excluidos): {len(abordados)}")
        print(f"  Leads novos restantes: {len(leads_filtrados)}")
        rodada_atual = db.get_rodada_atual(conn) + 1

    # Ordena
    leads_ordenados = ordenar_leads(leads_filtrados)

    # Limita
    leads_limitados = leads_ordenados[:args.limite]
    print(f"  Leads selecionados (limite {args.limite}): {len(leads_limitados)}")

    # Classifica nicho para filtro
    for lead in leads_limitados:
        lead["categoria"] = classificar_nicho(lead)

    # Mescla com status existente do banco
    for lead in leads_limitados:
        existing = db.get_lead(conn, lead["lead_id"])
        if existing:
            if existing["status"] != "novo":
                lead["status"] = existing["status"]
            if existing.get("data_abordagem"):
                lead["data_abordagem"] = existing["data_abordagem"]
            if existing.get("data_followup"):
                lead["data_followup"] = existing["data_followup"]
            if existing.get("observacoes"):
                lead["observacoes"] = existing["observacoes"]

    # Inserir leads no banco
    for lead in leads_limitados:
        db.upsert_lead(conn,
            lead_id=lead["lead_id"],
            nome=lead["nome"],
            status=lead.get("status", "novo"),
            data_abordagem=lead.get("data_abordagem", ""),
            data_followup=lead.get("data_followup", ""),
            observacoes=lead.get("observacoes", ""),
            nicho=lead.get("nicho", ""),
            cidade=lead.get("cidade", ""),
            telefone=lead.get("telefone", ""),
            whatsapp=lead.get("whatsapp", ""),
            score=lead.get("score", 0),
            prioridade=lead.get("prioridade", ""),
            mensagem_whatsapp=lead.get("mensagem_whatsapp", ""),
            link_whatsapp=lead.get("link_whatsapp", ""),
            rodada=rodada_atual)

    conn.commit()

    # Recalcular metricas
    db.recalcular_metricas(conn)
    metricas = db.get_metricas(conn)
    conn.close()

    # Cria diretório
    PAINEL_DIR.mkdir(parents=True, exist_ok=True)

    print(f"\n  Banco de dados: {db.DB_PATH}")
    print(f"  Rodada: {rodada_atual}")
    print(f"  Total de leads ja abordados (todas as rodadas): {metricas['total_abordados']}")
    print(f"  Mensagens enviadas (acumulado): {metricas['mensagens_enviadas']}")

    # Gera HTML
    html = gerar_html(leads_limitados, date.today(), metricas=metricas)
    caminho_html = PAINEL_DIR / "index.html"
    caminho_html.write_text(html, encoding="utf-8")
    print(f"  Painel salvo: {caminho_html}")
    print(f"  Tamanho: {len(html) // 1024} KB")

    # Resumo
    print(f"\n{'='*60}")
    print(f"  Painel gerado com sucesso!")
    print(f"  Leads: {len(leads_limitados)}")
    print(f"  Abra no navegador: {caminho_html}")
    print(f"\n  ROTINA RECOMENDADA:")
    print(f"  - Abra o painel no navegador")
    print(f"  - Copie a mensagem de cada lead")
    print(f"  - Clique em 'Abrir WhatsApp' para enviar")
    print(f"  - Marque o status após cada contato")
    print(f"  - Faça follow-up depois de 2 dias")
    print(f"  - NÃO envie mensagens em massa")
    print(f"  - Personalize quando possível")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    main()